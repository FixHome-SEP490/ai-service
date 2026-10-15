"""The knowledge base, laid out for a technician to check, and their answers taken back.

    python tools/kb_review.py confusions             # pairs the search mixes up with no question to tell them apart
    python tools/kb_review.py export --out kb-review.xlsx
    python tools/kb_review.py apply --file kb-review.xlsx [--write]

`export` writes one workbook with two sheets. "Lỗi" has every fault: the
appliance, the name, the symptoms customers use, the price range, how urgent,
what the customer is told to do. "Câu hỏi phân biệt" has every question the
assistant may ask back, the ones in use and the drafts in
app/data/discriminators_draft.json, with what a yes and a no point to. Each
row ends with "Duyệt" (Đúng / Sửa / Bỏ), "Sửa lại thành" and "Ghi chú" for the
technician.

`apply` reads the sheet back. A draft marked Đúng, or Sửa with the new
question written in, moves into the live table in fault_knowledge_base.json
and out of the drafts; nothing else is changed by the tool. Every other
answer (a fault marked Sửa, a live question marked Bỏ) is printed for a
person to make, because a symptom list or a price is not something to rewrite
from one cell. Without --write it only reports.

openpyxl is needed for export and apply only (it is not a service dependency).
"""

from __future__ import annotations

import argparse
import itertools
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
KB_PATH = ROOT / "app" / "data" / "fault_knowledge_base.json"
DRAFTS_PATH = ROOT / "app" / "data" / "discriminators_draft.json"

VERDICTS = ("Đúng", "Sửa", "Bỏ")
FAULT_HEADERS = [
    "Thiết bị", "Mã lỗi", "Tên lỗi", "Khách thường tả", "Giá từ", "Giá đến",
    "Mức khẩn", "Khuyên khách", "Duyệt", "Sửa lại thành", "Ghi chú",
]
QUESTION_HEADERS = [
    "Mã", "Trạng thái", "Thiết bị", "Câu hỏi", "Trả lời Có thì nghiêng về",
    "Trả lời Không thì nghiêng về", "Duyệt", "Sửa lại thành", "Ghi chú",
]
IN_USE = "Đang dùng"
DRAFT = "Nháp, chờ thợ duyệt"


def load() -> Tuple[dict, dict]:
    return (
        json.loads(KB_PATH.read_text(encoding="utf-8")),
        json.loads(DRAFTS_PATH.read_text(encoding="utf-8")),
    )


def _names(raw: dict) -> Dict[str, str]:
    return {f["fault_code"]: f["name_vi"] for f in raw["faults"]}


def separated_pairs(discriminators: Sequence[dict]) -> set:
    pairs = set()
    for d in discriminators:
        for a, b in itertools.product(d["favours_if_yes"], d["favours_if_no"]):
            pairs.add(frozenset((a, b)))
    return pairs


def confusions(raw: dict, retriever, discriminators: Sequence[dict]) -> List[Tuple[int, str, str, str]]:
    """Fault pairs the search ranks together, with no question that tells them apart.

    Each symptom a fault lists is searched as the customer would type it. A
    pair counts when the other fault comes first, or second within a fifth of
    the score. Most mixed-up pairs first.
    """
    known = separated_pairs(discriminators)
    counts: Counter = Counter()
    device_of = {f["fault_code"]: f["device_type"] for f in raw["faults"]}
    for fault in raw["faults"]:
        for symptom in fault["symptoms_vi"]:
            top = retriever.candidate_faults(symptom, fault["device_type"], top_k=3)
            if not top:
                continue
            codes = [t.fault.fault_code for t in top]
            if codes[0] != fault["fault_code"]:
                counts[frozenset((codes[0], fault["fault_code"]))] += 1
            elif len(top) > 1 and top[1].score >= 0.8 * top[0].score:
                counts[frozenset(codes[:2])] += 1
    rows = []
    for pair, count in counts.items():
        if len(pair) != 2 or pair in known:
            continue
        a, b = sorted(pair)
        rows.append((count, device_of[a], a, b))
    rows.sort(key=lambda r: (-r[0], r[1], r[2], r[3]))
    return rows


def fault_rows(raw: dict) -> List[List]:
    rows = []
    for f in sorted(raw["faults"], key=lambda f: (f["device_type"], f["fault_code"])):
        rows.append([
            f["device_type"], f["fault_code"], f["name_vi"], "; ".join(f.get("symptoms_vi", [])),
            f.get("price_min"), f.get("price_max"), f.get("urgency", ""),
            "; ".join(f.get("suggested_actions_vi", [])), "", "", "",
        ])
    return rows


def question_rows(raw: dict, drafts: dict) -> List[List]:
    names = _names(raw)

    def said(codes: Sequence[str]) -> str:
        return "; ".join(f"{names.get(c, c)} ({c})" for c in codes)

    rows = []
    for i, d in enumerate(raw["discriminators"], 1):
        rows.append([f"Q{i:02d}", IN_USE, d["device_type"], d["question_vi"], said(d["favours_if_yes"]), said(d["favours_if_no"]), "", "", ""])
    for d in drafts["drafts"]:
        rows.append([d["id"], DRAFT, d["device_type"], d["question_vi"], said(d["favours_if_yes"]), said(d["favours_if_no"]), "", "", ""])
    return rows


def apply_review(raw: dict, drafts: dict, reviewed: Sequence[Sequence]) -> Tuple[dict, dict, List[str]]:
    """Move approved drafts into the live table; report every other answer.

    `reviewed` is the question sheet as rows (headers dropped), in
    QUESTION_HEADERS order. Returns new copies; the inputs are not changed.
    """
    raw = json.loads(json.dumps(raw))
    drafts = json.loads(json.dumps(drafts))
    by_id = {d["id"]: d for d in drafts["drafts"]}
    report: List[str] = []
    moved: List[str] = []
    for row in reviewed:
        row = list(row) + [""] * (len(QUESTION_HEADERS) - len(row))
        code, status, _device, question, _yes, _no, verdict, rewrite, note = [
            (str(c).strip() if c is not None else "") for c in row[: len(QUESTION_HEADERS)]
        ]
        if not verdict:
            continue
        if verdict not in VERDICTS:
            report.append(f"{code}: không hiểu ô Duyệt '{verdict}', bỏ qua.")
            continue
        if status == DRAFT and code in by_id:
            draft = by_id[code]
            if verdict == "Bỏ":
                report.append(f"{code}: thợ bỏ câu nháp, giữ trong bản nháp để xem lại. {note}".strip())
                continue
            text = rewrite if verdict == "Sửa" else draft["question_vi"]
            if verdict == "Sửa" and not rewrite:
                report.append(f"{code}: đánh Sửa nhưng chưa ghi câu mới, chưa chuyển.")
                continue
            raw["discriminators"].append({
                "device_type": draft["device_type"],
                "question_vi": text,
                "favours_if_yes": list(draft["favours_if_yes"]),
                "favours_if_no": list(draft["favours_if_no"]),
            })
            moved.append(code)
            report.append(f"{code}: đã chuyển vào bảng đang dùng: {text}")
        elif verdict != "Đúng":
            report.append(f"{code} ({status}): thợ đánh {verdict}. {('Sửa thành: ' + rewrite + '. ') if rewrite else ''}{note}".strip() + " Cần người sửa tay.")
    drafts["drafts"] = [d for d in drafts["drafts"] if d["id"] not in moved]
    return raw, drafts, report


def fault_report(reviewed: Sequence[Sequence]) -> List[str]:
    """What the technician said about faults, for a person to act on."""
    lines = []
    for row in reviewed:
        row = list(row) + [""] * (len(FAULT_HEADERS) - len(row))
        cells = [(str(c).strip() if c is not None else "") for c in row[: len(FAULT_HEADERS)]]
        verdict, rewrite, note = cells[8], cells[9], cells[10]
        if verdict and verdict != "Đúng":
            lines.append(f"{cells[1]} {cells[2]}: thợ đánh {verdict}. {('Sửa thành: ' + rewrite + '. ') if rewrite else ''}{note}".strip())
    return lines


# ---------------------------------------------------------------- workbook I/O


def _sheet(wb, title: str, headers: List[str], rows: List[List], verdict_col: int, widths: Dict[int, int]):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    thin = Side(style="thin", color="BFBFBF")
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.row == 1)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="E8EEF7")
    for col, width in widths.items():
        ws.column_dimensions[chr(64 + col)].width = width
    ws.freeze_panes = "A2"
    letter = chr(64 + verdict_col)
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{len(rows) + 1}")
    return ws


def export(out: Path, raw: dict, drafts: dict) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Lỗi", FAULT_HEADERS, fault_rows(raw), 9, {1: 16, 2: 24, 3: 30, 4: 60, 5: 10, 6: 10, 7: 9, 8: 40, 9: 9, 10: 30, 11: 30})
    _sheet(wb, "Câu hỏi phân biệt", QUESTION_HEADERS, question_rows(raw, drafts), 7, {1: 6, 2: 18, 3: 16, 4: 50, 5: 40, 6: 40, 7: 9, 8: 40, 9: 30})
    wb.save(out)


def read_sheets(path: Path) -> Tuple[List[List], List[List]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    faults = [list(r) for r in wb["Lỗi"].iter_rows(min_row=2, values_only=True)]
    questions = [list(r) for r in wb["Câu hỏi phân biệt"].iter_rows(min_row=2, values_only=True)]
    return faults, questions


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)
    sub.add_parser("confusions")
    p_export = sub.add_parser("export")
    p_export.add_argument("--out", required=True)
    p_apply = sub.add_parser("apply")
    p_apply.add_argument("--file", required=True)
    p_apply.add_argument("--write", action="store_true")
    args = parser.parse_args()
    raw, drafts = load()

    if args.cmd == "confusions":
        sys.path.insert(0, str(ROOT))
        from app.services.pipeline.knowledge_base import get_knowledge_base
        from app.services.pipeline.retriever import Retriever

        names = _names(raw)
        draft_pairs = separated_pairs(drafts["drafts"])
        rows = confusions(raw, Retriever(get_knowledge_base()), raw["discriminators"])
        print(f"{len(rows)} cặp lỗi hay bị xếp sát nhau mà chưa có câu hỏi phân biệt:")
        for count, device, a, b in rows:
            mark = "  (đã có nháp)" if frozenset((a, b)) in draft_pairs else ""
            print(f"{count:3d}  {device:16s} {names[a]}  |  {names[b]}{mark}")
    elif args.cmd == "export":
        export(Path(args.out), raw, drafts)
        print(f"{len(raw['faults'])} lỗi, {len(raw['discriminators'])} câu hỏi đang dùng, {len(drafts['drafts'])} câu nháp -> {args.out}")
    else:
        fault_rows_back, question_rows_back = read_sheets(Path(args.file))
        new_raw, new_drafts, report = apply_review(raw, drafts, question_rows_back)
        for line in report + fault_report(fault_rows_back):
            print(line)
        if args.write:
            KB_PATH.write_text(json.dumps(new_raw, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            DRAFTS_PATH.write_text(json.dumps(new_drafts, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            print("Đã ghi fault_knowledge_base.json và discriminators_draft.json. Chạy lại pytest trước khi commit.")
        else:
            print("Chưa ghi gì (thêm --write để ghi).")


if __name__ == "__main__":
    main()
